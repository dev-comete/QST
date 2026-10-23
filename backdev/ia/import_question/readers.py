"""
LAYER 1 - READERS: uploaded file -> SourceDocument.

Goal: keep the signals the AI needs to tell right from wrong answers.
Highlighted cells / runs get the marker `[HL]`, bold text becomes **bold**.
The AI decides what the markers mean (see prompts.py).

PDFs and images are not converted: Gemini reads them natively (colors, layout, scans).

Each reader has the same signature `(filename, raw_bytes) -> SourceDocument`
so it can be tested without Django uploads.
"""
import io
import logging
import os

from .config import conf
from .exceptions import (
    DocumentTooLarge,
    EmptyDocument,
    QuestionImportError,
    UnreadableFile,
    UnsupportedFileType,
)
from .schemas import SourceDocument

logger = logging.getLogger(__name__)

MARK = "[HL]"


# ---------------------------------------------------------------------------
# TXT / MD / CSV
# ---------------------------------------------------------------------------
def _read_txt(filename: str, raw: bytes) -> SourceDocument:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return SourceDocument(filename=filename, text=raw.decode(encoding))
        except UnicodeDecodeError:
            continue
    raise UnreadableFile(detail="not utf-8 nor cp1252")


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------
_NO_COLOR_RGB = {"00000000", "FFFFFFFF"}


def _is_highlighted(cell) -> bool:
    """True if the cell has a solid, non-white background (any color: green, yellow...)."""
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return False
    color = fill.fgColor
    if color.type == "rgb":
        return color.rgb not in _NO_COLOR_RGB
    if color.type == "theme":
        return color.theme != 0          # theme 0 = white
    if color.type == "indexed":
        return color.indexed not in (64, 65)   # system fg/bg
    return False


def _looks_like_question_sheet(rows) -> bool:
    question_marks = sum(
        1 for row in rows for c in row if isinstance(c.value, str) and "?" in c.value
    )
    return question_marks >= 2


def _render_cell_value(value) -> str | None:
    if value is None:
        return None
    if hasattr(value, "isoformat"):          # dates are noise for question parsing
        return None
    text = str(value).replace("\n", " ").strip()
    return text or None


def _read_xlsx(filename: str, raw: bytes) -> SourceDocument:
    import openpyxl

    # NOT read_only: fills are not available in read-only mode.
    # data_only=True: use cached formula results instead of "=COUNTA(...)" strings.
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)

    sheets = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows()]
        sheets.append((ws.title, rows))

    selected = [(t, r) for t, r in sheets if _looks_like_question_sheet(r)]
    if not selected:                        # fallback: let the AI look at everything
        selected = sheets
    logger.info(
        "xlsx %s: %d sheets, kept %s",
        filename, len(sheets), [t for t, _ in selected],
    )

    lines: list[str] = []
    for title, rows in selected:
        lines.append(f"### SHEET: {title}")
        for row in rows:
            parts = []
            for cell in row:
                value = _render_cell_value(cell.value)
                if value is None:
                    continue
                prefix = f"{MARK} " if _is_highlighted(cell) else ""
                parts.append(f"{cell.column_letter}: {prefix}{value}")
            if parts:
                lines.append(f"R{row[0].row} | " + " | ".join(parts))
    return SourceDocument(filename=filename, text="\n".join(lines))


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def _render_paragraph(paragraph) -> str:
    parts: list[str] = []
    bold_buffer: list[str] = []
    highlighted = False

    def flush_bold():
        if bold_buffer:
            chunk = "".join(bold_buffer)
            parts.append(f"**{chunk.strip()}**" + (" " if chunk.endswith(" ") else ""))
            bold_buffer.clear()

    for run in paragraph.runs:
        if not run.text:
            continue
        if run.font.highlight_color is not None:
            highlighted = True
        if run.bold:
            bold_buffer.append(run.text)
        else:
            flush_bold()
            parts.append(run.text)
    flush_bold()

    text = "".join(parts).strip()
    if not text:
        return ""
    return f"{MARK} {text}" if highlighted else text


def _cell_is_shaded(cell) -> bool:
    from docx.oxml.ns import qn

    tc_pr = cell._tc.tcPr
    if tc_pr is None:
        return False
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        return False
    fill = shd.get(qn("w:fill"))
    return fill not in (None, "auto", "FFFFFF", "ffffff")


def _read_docx(filename: str, raw: bytes) -> SourceDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(raw))
    lines: list[str] = []

    # Walk the body in real order so tables stay between their paragraphs.
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            line = _render_paragraph(Paragraph(child, document))
            if line:
                lines.append(line)
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                seen, cells = set(), []
                for cell in row.cells:
                    if id(cell._tc) in seen:         # merged cells repeat
                        continue
                    seen.add(id(cell._tc))
                    text = " ".join(
                        t for t in (_render_paragraph(p) for p in cell.paragraphs) if t
                    )
                    if text and _cell_is_shaded(cell) and not text.startswith(MARK):
                        text = f"{MARK} {text}"
                    cells.append(text)
                if any(cells):
                    lines.append(" | ".join(cells))
    return SourceDocument(filename=filename, text="\n".join(lines))


# ---------------------------------------------------------------------------
# PDF / images: pass through, Gemini reads them natively
# ---------------------------------------------------------------------------
def _passthrough(mime_type: str):
    def reader(filename: str, raw: bytes) -> SourceDocument:
        return SourceDocument(filename=filename, data=raw, mime_type=mime_type)
    return reader


# ---------------------------------------------------------------------------
# Registry: add a new format = add one line here
# ---------------------------------------------------------------------------
_READERS = {
    ".txt": _read_txt,
    ".md": _read_txt,
    ".csv": _read_txt,
    ".xlsx": _read_xlsx,
    ".docx": _read_docx,
    ".pdf": _passthrough("application/pdf"),
    ".png": _passthrough("image/png"),
    ".jpg": _passthrough("image/jpeg"),
    ".jpeg": _passthrough("image/jpeg"),
    ".webp": _passthrough("image/webp"),
}

SUPPORTED_EXTENSIONS = tuple(_READERS)


# ---------------------------------------------------------------------------
# Public entry points
# ---------------------------------------------------------------------------
def read_upload(upload) -> SourceDocument:
    """Django UploadedFile -> SourceDocument."""
    filename = upload.name or ""
    extension = os.path.splitext(filename)[1].lower()
    reader = _READERS.get(extension)
    if reader is None:
        raise UnsupportedFileType(
            f"Format non supporté ({extension or 'sans extension'}). "
            f"Formats acceptés : {', '.join(SUPPORTED_EXTENSIONS)}."
        )

    limit = conf().max_file_bytes
    if upload.size > limit:
        raise DocumentTooLarge(f"Fichier trop volumineux (max {limit // (1024 * 1024)} Mo).")

    raw = upload.read()
    try:
        document = reader(filename, raw)
    except QuestionImportError:
        raise
    except Exception as exc:
        raise UnreadableFile(detail=repr(exc)) from exc

    _check_text_size(document)
    logger.info(
        "read %s (%s, %d bytes) -> %s",
        filename, extension, len(raw), "binary" if document.is_binary else f"{len(document.text)} chars",
    )
    logger.debug("annotated text sent to the AI:\n%s", document.text)
    return document


def read_raw_text(text: str) -> SourceDocument:
    """Text pasted by the user."""
    document = SourceDocument(filename="pasted_text", text=text)
    _check_text_size(document)
    return document


def _check_text_size(document: SourceDocument) -> None:
    if document.is_binary:
        return
    size = len((document.text or "").strip())
    cfg = conf()
    if size < cfg.min_text_chars:
        raise EmptyDocument()
    if size > cfg.max_text_chars:
        raise DocumentTooLarge(
            f"Texte trop long ({size} caractères, max {cfg.max_text_chars}). Découpez le document."
        )
