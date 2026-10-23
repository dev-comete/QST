"""
One test class per layer. No database, no network (the LLM is mocked).
Run with:  python manage.py test ia.import_questions
"""
import io
from types import SimpleNamespace
from unittest import mock

import docx
import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from openpyxl.styles import PatternFill

from . import llm, presenter, readers, service
from .exceptions import (
    EmptyDocument, LLMOutputInvalid, LLMTruncated, UnsupportedFileType,
)
from .normalizer import normalize
from .schemas import AIExtraction, AIOption, AIQuestion

GREEN = PatternFill("solid", fgColor="FF6AA84F")


def _upload(name, raw):
    return SimpleUploadedFile(name, raw)


def _xlsx_bytes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quiz"
    ws["A1"], ws["B1"] = "1- Capital de la France ?", "Berlin"
    ws["B2"] = "Paris"
    ws["B2"].fill = GREEN
    ws["A4"], ws["B4"] = "2- Combien font 2+2 ?", "4"
    ws["B4"].fill = GREEN
    noise = wb.create_sheet("Suivi")
    noise["A1"] = "Nom"
    noise["A2"] = "Dupont"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ReadersTests(SimpleTestCase):
    def test_xlsx_marks_highlighted_cells_and_skips_noise_sheet(self):
        doc = readers.read_upload(_upload("q.xlsx", _xlsx_bytes()))
        self.assertIn("B: [HL] Paris", doc.text)
        self.assertIn("B: Berlin", doc.text)
        self.assertNotIn("Dupont", doc.text)

    def test_docx_marks_bold_and_highlight(self):
        d = docx.Document()
        p = d.add_paragraph()
        p.add_run("Bonne réponse").bold = True
        p2 = d.add_paragraph()
        r = p2.add_run("Autre")
        r.font.highlight_color = 7
        buf = io.BytesIO()
        d.save(buf)
        doc = readers.read_upload(_upload("q.docx", buf.getvalue()))
        self.assertIn("**Bonne réponse**", doc.text)
        self.assertIn("[HL] Autre", doc.text)

    def test_txt_falls_back_to_cp1252(self):
        raw = "Quelle est la réponse à cette question ?".encode("cp1252")
        doc = readers.read_upload(_upload("q.txt", raw))
        self.assertIn("réponse", doc.text)

    def test_pdf_is_passthrough(self):
        doc = readers.read_upload(_upload("q.pdf", b"%PDF-1.4 fake"))
        self.assertTrue(doc.is_binary)
        self.assertEqual(doc.mime_type, "application/pdf")

    def test_unsupported_extension(self):
        with self.assertRaises(UnsupportedFileType):
            readers.read_upload(_upload("q.exe", b"x"))

    def test_too_short_text(self):
        with self.assertRaises(EmptyDocument):
            readers.read_raw_text("hi")


class NormalizerTests(SimpleTestCase):
    def _extract(self, *questions):
        return AIExtraction(questions=list(questions))

    def test_type_inference(self):
        qcu = AIQuestion(question="1- Q ?", options=[AIOption(text="a", is_correct=True), AIOption(text="b", is_correct=False)])
        qcm = AIQuestion(question="Q2 ?", options=[AIOption(text="a", is_correct=True), AIOption(text="b", is_correct=True)])
        ouv = AIQuestion(question="Q3 ?", options=[])
        result = normalize(self._extract(qcu, qcm, ouv))
        self.assertEqual([q.type_code for q in result.questions], ["QCU", "QCM", "OUV"])
        self.assertEqual(result.questions[0].question, "Q ?")      # numbering stripped

    def test_warnings(self):
        none_ok = AIQuestion(question="Q ?", options=[AIOption(text="a", is_correct=False), AIOption(text="b", is_correct=False)])
        single = AIQuestion(question="Q ?", options=[AIOption(text="a", is_correct=True)])
        result = normalize(self._extract(none_ok, single))
        self.assertEqual(result.questions[0].warnings, ["no_correct_answer"])
        self.assertEqual(result.questions[1].warnings, ["too_few_options"])

    def test_cleaning_dedupe_and_skip(self):
        q = AIQuestion(question="Q ?", options=[
            AIOption(text="A) [HL] Paris", is_correct=True),
            AIOption(text="paris", is_correct=False),
            AIOption(text="   ", is_correct=False),
        ])
        empty = AIQuestion(question="   ", options=[])
        result = normalize(self._extract(q, empty))
        self.assertEqual(result.skipped, 1)
        self.assertEqual(len(result.questions[0].options), 1)
        self.assertEqual(result.questions[0].options[0].text, "Paris")
        self.assertTrue(result.questions[0].options[0].is_correct)


class LLMTests(SimpleTestCase):
    def test_truncated(self):
        resp = SimpleNamespace(candidates=[SimpleNamespace(finish_reason="FinishReason.MAX_TOKENS")], parsed=None, text=None)
        with self.assertRaises(LLMTruncated):
            llm._parse_response(resp)

    def test_parsed_ok(self):
        parsed = AIExtraction(questions=[])
        resp = SimpleNamespace(candidates=[SimpleNamespace(finish_reason="STOP")], parsed=parsed, text="")
        self.assertIs(llm._parse_response(resp), parsed)

    def test_empty_response(self):
        resp = SimpleNamespace(candidates=[], parsed=None, text=None)
        with self.assertRaises(LLMOutputInvalid):
            llm._parse_response(resp)

    def test_text_fallback(self):
        resp = SimpleNamespace(candidates=[SimpleNamespace(finish_reason="STOP")], parsed=None, text='{"questions": []}')
        self.assertEqual(llm._parse_response(resp).questions, [])


class PresenterTests(SimpleTestCase):
    def test_unknown_type_is_flagged_not_dropped(self):
        q = AIQuestion(question="Q ?", options=[AIOption(text="a", is_correct=True)])
        payload = presenter.to_api_payload(normalize(AIExtraction(questions=[q])), {"QCM": 2})
        self.assertEqual(payload["count"], 1)
        self.assertIsNone(payload["questions"][0]["type_id"])
        self.assertIn("unknown_type", payload["questions"][0]["warnings"])


class ServiceTests(SimpleTestCase):
    def test_pipeline_with_mocked_llm(self):
        fake = AIExtraction(questions=[AIQuestion(question="Capitale ?", options=[
            AIOption(text="Paris", is_correct=True), AIOption(text="Berlin", is_correct=False)])])
        with mock.patch.object(service, "extract_questions", return_value=fake) as m, \
             mock.patch.object(service.cache, "get", return_value=None), \
             mock.patch.object(service.cache, "set"):
            result = service.import_questions(upload=_upload("q.xlsx", _xlsx_bytes()))
        m.assert_called_once()
        self.assertEqual(result.questions[0].type_code, "QCU")
